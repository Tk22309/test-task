import os, json, re
from pathlib import Path
from typing import Optional
import time
from openai import OpenAI
from openpyxl import load_workbook, Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont

API_KEY = "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"
client = OpenAI(api_key=API_KEY)


def Give_score(text: str, max_retries: int = 3, delay: float = 5.0) -> str:
    messages = [#повыдомлення для чату джпт для отримання оцінки
        {
            "role": "system",
            "content": (
                "Ти — аналітик якості сервісу автосервісу, який оцінює роботу менеджера, текст може містити помилки тому ураховй цу при оцнці "
                "за визначеними критеріями. Ти повертаєш лише JSON, без зайвого тексту."
            ),
        },
        {
            "role": "user",
            "content": f"""
Оціни роботу менеджера за записом розмови текстом:

{text}

Відповідь надай у форматі JSON з оцінками за критеріями (може бути кілька значень):
{{
"Дата":"-",
"Тип звернення":"консультація",
"Номер телефону":"-",
"Філія":"-",
"Менеджер":"знайди імя якщо вказано то ще фамілію  менеджера якщо не знайшов то наиши(-)",
"Початок розмови, представлення":"так або ні",
"Чи дізнався менеджер кузов автомобіля":"якщо так то напиши (серію марку та модель машини), якщо ні то слово (ні)",
"Чи дізнався менеджер рік автомобіля":"(рік)/(ні)",
"Чи дізнався менеджер пробіг":"(пробіг)/(ні)",
"Пропозиція про комплексну діагностику":"так або ні",
"Дізнався які роботи робилися раніше":"(які роботи велись)/(ні)",
"Запис на сервіс, Дата":"(день тижня (дата) час)",
"Завершення розмови прощання":"так або ні",
"Яка робота з топ 100":"-",
"Чи дотримувався всіх інструкцій з топ 100 робіт":"так або ні",
"Яких рекомендацій менеджер не дотримувався з топ 100 робіт":"-",
"Результат":"число 0-100",
"Запчастини":"що цікавить клієнта і які деталі потрібні",
"Коментар":"які фрази були непідходящі та що не забувати менеджеру",
"над якими фразами менеджер має попрацювати":"візьми ЦІЛІ речення з тексту без змін; якщо кілька — розділи ';'"
}}
""",
        },
    ]

    for attempt in range(1, max_retries + 1):
        try:
            print(f"[i] GPT-запит (спроба {attempt}/{max_retries})...")
            resp = client.chat.completions.create(
                model="gpt-4o-mini",
                messages=messages,
                temperature=0.0,
                timeout=90  # час очікування в секундах
            )
            raw = resp.choices[0].message.content or ""
            clean_json = re.sub(r"^```(?:json)?\s*|\s*```$", "", raw.strip(), flags=re.DOTALL)
            print(f"[✓] Отримано відповідь ({len(clean_json)} символів)")
            return clean_json

        except Exception as e:
            print(f"[!] Помилка GPT: {e}")
            if attempt < max_retries:
                print(f"[→] Повтор через {delay} секунд...")
                time.sleep(delay)
            else:
                print("[×] Досягнуто межі повторів. Пропускаємо цей файл.")
                return json.dumps({"Помилка": f"Не вдалося отримати відповідь GPT: {str(e)}"})


#Запис JSON у Excel
def write_json_to_xlsx(
    xlsx_path: str,
    clean_json_text: str,
    sheet_name: Optional[str] = None,
    header_row: int = 2,
    full_call_text: Optional[str] = None,
) -> int:
    def _normalize_key(k: str) -> str:
        k = str(k).strip()
        aliases = {
            "Чи дізнвся менеджер кузов атвомобіля": "Чи дізнався менеджер кузов автомобіля",
            "Чи дізнався менеджр пробіг": "Чи дізнався менеджер пробіг",
            "Яких рекоменадцій менеджер не дотримувався з топ 100 робіт": "Яких рекомендацій менеджер не дотримувався з топ 100 робіт",
            "тип звернення": "Тип звернення",
            "дата": "Дата",
        }
        return aliases.get(k, k)

    def _build_rich_text(full_text: str, red_ranges: list[tuple[int, int]]):
        rt = CellRichText()
        normal = InlineFont()
        red = InlineFont(color="FF0000")
        pos = 0
        for s, e in red_ranges:
            if pos < s:
                rt.append(TextBlock(text=full_text[pos:s], font=normal))
            rt.append(TextBlock(text=full_text[s:e], font=red))
            pos = e
        if pos < len(full_text):
            rt.append(TextBlock(text=full_text[pos:], font=normal))
        return rt

    try:
        pairs = json.loads(clean_json_text, object_pairs_hook=list)
    except json.JSONDecodeError:
        pairs = json.loads(clean_json_text.replace("'", '"'), object_pairs_hook=list)

    occurrences: dict[str, list[str]] = {}
    for k, v in pairs:
        nk = _normalize_key(k)
        s = "" if v is None else str(v)
        s = re.sub(rf"^\s*{re.escape(nk)}\s*:\s*", "", s).strip()
        occurrences.setdefault(nk, []).append(s)

    def first(key: str) -> str:
        for x in occurrences.get(key, []):
            if str(x).strip():
                return str(x)
        return ""

    def second(key: str) -> str:
        arr = [str(x) for x in occurrences.get(key, []) if str(x).strip()]
        return arr[1] if len(arr) >= 2 else ""

    col_values = {
        "A": full_call_text or "",
        "B": first("Тип звернення"),
        "C": first("Номер телефону"),
        "D": first("Філія"),
        "E": first("Менеджер"),
        "F": first("Початок розмови, представлення"),
        "G": first("Чи дізнався менеджер кузов автомобіля"),
        "H": first("Чи дізнався менеджер рік автомобіля"),
        "I": first("Чи дізнався менеджер пробіг"),
        "J": first("Пропозиція про комплексну діагностику"),
        "K": first("Дізнався які роботи робилися раніше"),
        "L": first("Запис на сервіс, Дата"),
        "M": first("Завершення розмови прощання"),
        "N": first("Яка робота з топ 100"),
        "O": first("Чи дотримувався всіх інструкцій з топ 100 робіт"),
        "P": first("Яких рекомендацій менеджер не дотримувався з топ 100 робіт"),
        "Q": first("Результат"),
        "R": second("Дата"),
        "S": second("Тип звернення"),
        "T": first("Запчастини"),
        "U": first("Коментар"),
    }

    xlsx = Path(xlsx_path)
    if not xlsx.exists():
        Workbook().save(xlsx) 

    wb = load_workbook(xlsx)
    ws = wb[sheet_name] if sheet_name else wb.active

    row = header_row + 1
    while ws.cell(row=row, column=1).value not in (None, ""):
        row += 1

    for letter, value in col_values.items():
        col_num = ord(letter) - ord("A") + 1
        ws.cell(row=row, column=col_num, value=value)

    # підкреслюємо проблемні фрази у колонці A
    full_text = col_values["A"]
    phrases_text = first("над якими фразами менеджер має попрацювати")
    phrases = [p.strip() for p in re.split(r"[;\n；]+", phrases_text or "") if p.strip()]

    if full_text and phrases:
        ranges: list[tuple[int, int]] = []
        for ph in phrases:
            start = 0
            while True:
                i = full_text.find(ph, start)
                if i == -1:
                    break
                ranges.append((i, i + len(ph)))
                start = i + len(ph)
        if ranges:
            ranges.sort()
            merged = [ranges[0]]
            for s, e in ranges[1:]:
                ls, le = merged[-1]
                if s <= le:
                    merged[-1] = (ls, max(le, e))
                else:
                    merged.append((s, e))
            ws.cell(row=row, column=1).value = _build_rich_text(full_text, merged)

    wb.save(xlsx_path)
    return row



